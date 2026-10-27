"""Conserva columnas/valores y consolida una fila por NumeroPedido.
Requiere openpyxl solo para regenerar los CSV, no para importar en Genesys.
"""
import argparse
import csv
import datetime
from pathlib import Path
import openpyxl


def normalize(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def load_rows(source):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    sheet = workbook.active
    if any(c.data_type == 'f' for row in sheet.iter_rows() for c in row):
        raise ValueError('Se detectaron formulas: resolver sus valores antes de preparar datos exactos.')
    iterator = sheet.iter_rows(values_only=True)
    columns = list(next(iterator))
    if columns[0] != 'NumeroPedido' or len(set(columns)) != len(columns):
        raise ValueError('Se esperan cabeceras unicas empezando por NumeroPedido.')
    rows = [dict(zip(columns, map(normalize, row))) for row in iterator if any(v is not None for v in row)]
    return columns, rows


def consolidate(columns, rows):
    groups = {}
    for row in rows:
        key = row['NumeroPedido']
        if not key:
            raise ValueError('Hay una fila sin NumeroPedido.')
        groups.setdefault(key, []).append(row)
    return [{name: ' | '.join(dict.fromkeys(row[name] for row in items)) for name in columns} for items in groups.values()]


def write_csv(path, columns, rows):
    # No sobrescribir los archivos de entrega o datos del usuario por accidente.
    with path.open('x', encoding='utf-8-sig', newline='') as output:
        writer = csv.DictWriter(output, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('output_directory', type=Path)
    args = parser.parse_args()
    columns, original = load_rows(args.source)
    consolidated = consolidate(columns, original)
    args.output_directory.mkdir(parents=True, exist_ok=True)
    for name in ['original.csv', 'consolidado.csv']:
        if (args.output_directory/name).exists():
            raise FileExistsError('Elegir una carpeta de salida nueva; no se sobrescribe.')
    write_csv(args.output_directory/'original.csv', columns, original)
    write_csv(args.output_directory/'consolidado.csv', columns, consolidated)
    print(f'{len(columns)} columnas; {len(original)} filas originales; {len(consolidated)} pedidos.')


if __name__ == '__main__':
    main()
